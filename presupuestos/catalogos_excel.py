# presupuestos/catalogos_excel.py
#
# Build/read/apply the Excel workbook used at /admin/catalogos/ to
# bulk-load Categoria, TipoGasto, CuentaContableTipoGasto, and
# CategoriaProductoTipoGasto. Odoo fetch helpers here are also imported by
# scripts/classify_odoo_catalog.py, so there's one source of truth for
# "which categories/accounts exist" between the Excel flow and the
# keyword-based classifier script.

from io import BytesIO

import openpyxl
from django.db import transaction
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from core.database.odoo import get_odoo_connection

from .models import Categoria, CategoriaProductoTipoGasto, CuentaContableTipoGasto, TipoGasto

MAX_FILAS_VALIDACION = 500

HEADER_FILL = PatternFill(start_color="035953", end_color="035953", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

HOJAS_REQUERIDAS = ["Categorias", "Tipos de gasto", "Categorias de producto", "Cuentas contables"]


class ErroresPlantilla(Exception):
    def __init__(self, errores):
        self.errores = errores
        super().__init__("; ".join(errores))


def fetch_odoo_product_categories():
    uid, models_proxy, db, password = get_odoo_connection()
    return models_proxy.execute_kw(
        db, uid, password, "product.category", "search_read", [[]], {"fields": ["id", "name"]}
    )


def fetch_odoo_direct_expense_accounts():
    """
    {odoo_account_id: odoo_account_name} for every non-GRNI account actually
    used on a paid/in_payment vendor-bill line - same source set the daily
    scheduler classifies against, so nothing shows up here that GastoReal
    couldn't already need mapped.
    """
    uid, models_proxy, db, password = get_odoo_connection()
    bill_ids = models_proxy.execute_kw(
        db, uid, password, "account.move", "search_read",
        [[["move_type", "=", "in_invoice"], ["payment_state", "in", ["paid", "in_payment"]]]],
        {"fields": ["id"], "limit": 20000},
    )
    ids = [b["id"] for b in bill_ids]

    accounts_seen = {}
    CHUNK = 500
    for i in range(0, len(ids), CHUNK):
        chunk_ids = ids[i:i + CHUNK]
        lines = models_proxy.execute_kw(
            db, uid, password, "account.move.line", "search_read",
            [[["move_id", "in", chunk_ids], ["display_type", "=", "product"]]],
            {"fields": ["account_id"]},
        )
        for line in lines:
            acc = line["account_id"]
            if acc and "Goods Received" not in acc[1]:
                accounts_seen[acc[0]] = acc[1]
    return accounts_seen


def _hoja_con_encabezado(wb, titulo, encabezados):
    ws = wb.create_sheet(titulo)
    ws.append(encabezados)
    for col in range(1, len(encabezados) + 1):
        celda = ws.cell(row=1, column=col)
        celda.font = HEADER_FONT
        celda.fill = HEADER_FILL
    return ws


def _autoajustar_columnas(ws):
    for columna in ws.columns:
        largo = max((len(str(c.value)) if c.value is not None else 0) for c in columna)
        ws.column_dimensions[columna[0].column_letter].width = min(max(largo + 2, 12), 60)


def construir_plantilla_excel():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_categorias = _hoja_con_encabezado(wb, "Categorias", ["Nombre"])
    for categoria in Categoria.objects.order_by("nombre"):
        ws_categorias.append([categoria.nombre])

    ws_tipos = _hoja_con_encabezado(wb, "Tipos de gasto", ["Nombre", "Categoria", "Descripcion"])
    for tipo in TipoGasto.objects.select_related("categoria").order_by("nombre"):
        ws_tipos.append([tipo.nombre, tipo.categoria.nombre, tipo.descripcion])

    dv_categoria = DataValidation(
        type="list", formula1=f"='Categorias'!$A$2:$A${MAX_FILAS_VALIDACION}", allow_blank=True
    )
    ws_tipos.add_data_validation(dv_categoria)
    dv_categoria.add(f"B2:B{MAX_FILAS_VALIDACION}")

    ws_cat_producto = _hoja_con_encabezado(
        wb, "Categorias de producto", ["Odoo Category ID", "Odoo Category Name", "Tipo de gasto"]
    )
    mapeo_producto_existente = {
        m.odoo_category_id: (m.tipo_gasto.nombre if m.tipo_gasto_id else "")
        for m in CategoriaProductoTipoGasto.objects.select_related("tipo_gasto")
    }
    for cat in fetch_odoo_product_categories():
        ws_cat_producto.append([cat["id"], cat["name"], mapeo_producto_existente.get(cat["id"], "")])

    ws_cuentas = _hoja_con_encabezado(
        wb, "Cuentas contables", ["Odoo Account ID", "Odoo Account Name", "Tipo de gasto"]
    )
    mapeo_cuenta_existente = {
        m.odoo_account_id: (m.tipo_gasto.nombre if m.tipo_gasto_id else "")
        for m in CuentaContableTipoGasto.objects.select_related("tipo_gasto")
    }
    for acc_id, acc_name in fetch_odoo_direct_expense_accounts().items():
        ws_cuentas.append([acc_id, acc_name, mapeo_cuenta_existente.get(acc_id, "")])

    for ws in (ws_cat_producto, ws_cuentas):
        dv_tipo_gasto = DataValidation(
            type="list", formula1=f"='Tipos de gasto'!$A$2:$A${MAX_FILAS_VALIDACION}", allow_blank=True
        )
        ws.add_data_validation(dv_tipo_gasto)
        dv_tipo_gasto.add(f"C2:C{MAX_FILAS_VALIDACION}")

    for ws in wb.worksheets:
        _autoajustar_columnas(ws)

    return wb


def plantilla_a_bytes():
    wb = construir_plantilla_excel()
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _filas_no_vacias(ws):
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and any(v not in (None, "") for v in row):
            yield row


def _texto(valor):
    return str(valor).strip() if valor not in (None, "") else ""


def procesar_plantilla_excel(archivo):
    """
    Parses and cross-validates the uploaded workbook. Raises ErroresPlantilla
    (never a partial result) if anything doesn't resolve - a Tipo de gasto
    naming a Categoria that isn't in the Categorias sheet, or a mapping row
    naming a Tipo de gasto that isn't in the Tipos de gasto sheet.
    """
    try:
        wb = openpyxl.load_workbook(archivo, data_only=True)
    except Exception as exc:
        raise ErroresPlantilla([f"No se pudo leer el archivo como Excel (.xlsx): {exc}"]) from exc

    errores = [f"Falta la hoja '{hoja}' en el archivo." for hoja in HOJAS_REQUERIDAS if hoja not in wb.sheetnames]
    if errores:
        raise ErroresPlantilla(errores)

    categorias_nombres = set()
    for row in _filas_no_vacias(wb["Categorias"]):
        nombre = _texto(row[0] if len(row) > 0 else None)
        if nombre:
            categorias_nombres.add(nombre)

    tipos_rows = []
    tipos_nombres = set()
    for row in _filas_no_vacias(wb["Tipos de gasto"]):
        nombre = _texto(row[0] if len(row) > 0 else None)
        categoria_nombre = _texto(row[1] if len(row) > 1 else None)
        descripcion = _texto(row[2] if len(row) > 2 else None)
        if not nombre:
            continue
        if not categoria_nombre:
            errores.append(f"Tipo de gasto '{nombre}': falta la Categoria.")
            continue
        if categoria_nombre not in categorias_nombres:
            errores.append(
                f"Tipo de gasto '{nombre}': la categoria '{categoria_nombre}' no esta en la hoja Categorias."
            )
            continue
        tipos_nombres.add(nombre)
        tipos_rows.append((nombre, categoria_nombre, descripcion))

    def leer_mapeo(ws, etiqueta):
        filas_out = []
        for row in _filas_no_vacias(ws):
            odoo_id_raw = row[0] if len(row) > 0 else None
            odoo_nombre = _texto(row[1] if len(row) > 1 else None)
            tipo_gasto_nombre = _texto(row[2] if len(row) > 2 else None)
            if odoo_id_raw is None or not odoo_nombre:
                continue
            try:
                odoo_id = int(odoo_id_raw)
            except (TypeError, ValueError):
                errores.append(f"{etiqueta} '{odoo_nombre}': el ID '{odoo_id_raw}' no es un numero valido.")
                continue
            if tipo_gasto_nombre and tipo_gasto_nombre not in tipos_nombres:
                errores.append(
                    f"{etiqueta} '{odoo_nombre}' (id {odoo_id}): el tipo de gasto '{tipo_gasto_nombre}' "
                    "no esta en la hoja Tipos de gasto."
                )
                continue
            filas_out.append((odoo_id, odoo_nombre, tipo_gasto_nombre or None))
        return filas_out

    cuentas_rows = leer_mapeo(wb["Cuentas contables"], "Cuenta contable")
    categorias_producto_rows = leer_mapeo(wb["Categorias de producto"], "Categoria de producto")

    if errores:
        raise ErroresPlantilla(errores)

    return {
        "categorias": sorted(categorias_nombres),
        "tipos": tipos_rows,
        "cuentas": cuentas_rows,
        "categorias_producto": categorias_producto_rows,
    }


def aplicar_plantilla(datos, borrar_todo):
    """
    borrar_todo=True wipes Categoria/TipoGasto/both mapping tables first
    (children before parents, so TipoGasto's own PROTECT constraints from
    the mapping tables are already cleared by the time it's deleted).
    TipoGasto is still PROTECTed by Presupuesto - if any exists, this raises
    django.db.models.ProtectedError and the whole transaction rolls back
    (the caller is expected to catch it and leave the flag untouched).
    borrar_todo=False only upserts: existing rows not mentioned in the
    workbook are left as-is.
    """
    resumen = {"categorias": 0, "tipos": 0, "cuentas": 0, "categorias_producto": 0}

    with transaction.atomic():
        if borrar_todo:
            CategoriaProductoTipoGasto.objects.all().delete()
            CuentaContableTipoGasto.objects.all().delete()
            TipoGasto.objects.all().delete()
            Categoria.objects.all().delete()

        categorias_obj = {}
        for nombre in datos["categorias"]:
            obj, _ = Categoria.objects.get_or_create(nombre=nombre)
            categorias_obj[nombre] = obj
            resumen["categorias"] += 1

        tipos_obj = {}
        for nombre, categoria_nombre, descripcion in datos["tipos"]:
            categoria = categorias_obj.get(categoria_nombre) or Categoria.objects.get(nombre=categoria_nombre)
            obj, _ = TipoGasto.objects.update_or_create(
                nombre=nombre, defaults={"categoria": categoria, "descripcion": descripcion},
            )
            tipos_obj[nombre] = obj
            resumen["tipos"] += 1

        for odoo_id, odoo_nombre, tipo_nombre in datos["cuentas"]:
            CuentaContableTipoGasto.objects.update_or_create(
                odoo_account_id=odoo_id,
                defaults={"odoo_account_name": odoo_nombre, "tipo_gasto": tipos_obj.get(tipo_nombre)},
            )
            resumen["cuentas"] += 1

        for odoo_id, odoo_nombre, tipo_nombre in datos["categorias_producto"]:
            CategoriaProductoTipoGasto.objects.update_or_create(
                odoo_category_id=odoo_id,
                defaults={"odoo_category_name": odoo_nombre, "tipo_gasto": tipos_obj.get(tipo_nombre)},
            )
            resumen["categorias_producto"] += 1

    return resumen
